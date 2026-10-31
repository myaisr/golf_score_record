import copy
import importlib.util
import json
from pathlib import Path
import shutil
import subprocess
import tempfile
import unittest
from datetime import datetime

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location('roundnote_excel', ROOT / 'scripts' / 'import_rounds.py')
excel = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(excel)


def make_round(identifier='round-one', day='2026-11-01'):
    return {'id': identifier, 'date': day, 'place': '東京ゴルフ場', 'active': 17,
            'finished': True, 'createdAt': day + 'T01:00:00.000Z', 'updatedAt': day + 'T02:00:00.000Z',
            'holes': [{'par': 4, 'long': 2, 'short': 1, 'putts': 2, 'obTee': 0,
                       'obLater': 0, 'direct': False, 'tee': 'straight', 'teeShort': False, 'approach': None}
                      for _ in range(18)]}


class ImportTests(unittest.TestCase):
    def test_tee_club_column_and_canonical_backup_preserve_all_four_choices(self):
        record = make_round()
        clubs = list(excel.CLUBS)
        for index, club in enumerate(clubs):
            record['holes'][index]['teeClub'] = club
        self.write_json([record])
        excel.import_file(self.source, self.output)
        workbook = load_workbook(self.output, data_only=True)
        for index, club in enumerate(clubs):
            self.assertEqual(workbook[excel.DETAIL].cell(index + 2, 30).value, excel.CLUBS[club])
        self.assertIsNone(workbook[excel.DETAIL].cell(6, 30).value)
        canonical = json.loads(workbook[excel.HISTORY].cell(2, excel.JSON_COLUMN).value)
        self.assertEqual(canonical['holes'][0]['teeClub'], 'driver')
        workbook.close()
        self.assertEqual(excel.import_file(self.source, self.output), 1)
        record['holes'][0]['teeClub'] = 'putter'
        with self.assertRaises(ValueError):
            excel.normalize_round(record)

    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix='roundnote-test-')
        self.directory = Path(self.temporary.name)
        self.source = self.directory / 'rounds.json'
        self.output = self.directory / 'rounds.xlsx'

    def tearDown(self):
        self.temporary.cleanup()

    def write_json(self, rounds):
        self.source.write_text(json.dumps({'schemaVersion': 2, 'currentRoundId': rounds[0]['id'], 'goals': excel.DEFAULT_GOALS, 'rounds': rounds}, ensure_ascii=False), encoding='utf-8')

    def test_upsert_and_older_data_cannot_duplicate_or_revert(self):
        record = make_round()
        self.write_json([record])
        self.assertEqual(excel.import_file(self.source, self.output), 1)
        self.assertEqual(excel.import_file(self.source, self.output), 1)
        updated = copy.deepcopy(record)
        updated['holes'][0]['long'] = 3
        updated['updatedAt'] = '2026-11-02T02:00:00.000Z'
        self.write_json([updated])
        excel.import_file(self.source, self.output)
        self.write_json([record])
        excel.import_file(self.source, self.output)
        workbook = load_workbook(self.output, data_only=True)
        self.assertEqual(workbook[excel.HISTORY].max_row, 2)
        self.assertEqual(workbook[excel.DETAIL].max_row, 19)
        self.assertEqual(workbook[excel.HISTORY]['H2'].value, 91)
        self.assertEqual(workbook[excel.DETAIL]['G2'].value, 3)
        workbook.close()

    def test_formulas_cached_values_charts_and_three_sheets(self):
        record = make_round()
        record['holes'][0].update(direct=True, short=0, obTee=2, tee='left', teeShort=True)
        record['holes'][1].update(short=2, obLater=1, tee='right')
        record['holes'][2].update(putts=3)
        self.write_json([record])
        excel.import_file(self.source, self.output)
        workbook = load_workbook(self.output, data_only=True)
        self.assertEqual(workbook.sheetnames, [excel.DASHBOARD, excel.HISTORY, excel.DETAIL])
        history = workbook[excel.HISTORY]
        self.assertEqual(history['H2'].value, 91)
        self.assertEqual(history['M2'].value, 3)
        self.assertEqual(history['P2'].value, 2)
        self.assertEqual(history['Q2'].value, 0)
        self.assertEqual(history['W2'].value, 17)
        self.assertEqual(history['X2'].value, 1)
        self.assertAlmostEqual(history['Y2'].value, 1 / 17)
        self.assertEqual(history['AD2'].value, 1)
        self.assertEqual(history['AB2'].value, 1)
        self.assertAlmostEqual(history['AC2'].value, 1 / 18)
        self.assertIsInstance(history['B2'].value, datetime)
        self.assertTrue(history.column_dimensions['AZ'].hidden)
        self.assertEqual(len(workbook[excel.DASHBOARD]._charts), 6)
        workbook.close()
        formulas = load_workbook(self.output, data_only=False)
        self.assertTrue(formulas[excel.HISTORY]['Y2'].value.startswith('=IF('))
        self.assertTrue(formulas[excel.DASHBOARD]['B9'].value.startswith('=IF('))
        self.assertIn("'分析ダッシュボード'!", formulas[excel.DASHBOARD]._charts[0].series[0].val.numRef.f)
        self.assertEqual(formulas[excel.DASHBOARD]._charts[1].series[0].title.v, 'ティーOB左')
        formulas.close()

    def test_no_eligible_approaches_are_blank_not_zero(self):
        record = make_round()
        for hole in record['holes']:
            hole.update(short=0, direct=True)
        self.write_json([record])
        excel.import_file(self.source, self.output)
        workbook = load_workbook(self.output, data_only=True)
        self.assertEqual(workbook[excel.HISTORY]['W2'].value, 0)
        self.assertIsNone(workbook[excel.HISTORY]['Y2'].value)
        self.assertIsNone(workbook[excel.DASHBOARD]['B9'].value)
        self.assertEqual(workbook[excel.HISTORY]['AC2'].value, 1)
        workbook.close()

    def test_legacy_unknown_short_flags_are_preserved_and_excluded_from_rate(self):
        legacy = make_round()
        for key in ('id', 'createdAt', 'updatedAt', 'finished'):
            legacy.pop(key)
        for hole in legacy['holes']:
            hole.pop('teeShort')
            hole['tee'] = 'good'
        record = excel.read_json(legacy)[0]
        metrics = excel.analyze(record['holes'])
        self.assertEqual(metrics['teeRecorded'], 18)
        self.assertEqual(metrics['teeShortRecorded'], 0)
        self.assertIsNone(metrics['teeShortRate'])
        self.assertTrue(all(h['teeShort'] is None for h in record['holes']))
        self.write_json([record])
        excel.import_file(self.source, self.output)
        workbook = load_workbook(self.output, data_only=True)
        self.assertEqual(workbook[excel.HISTORY]['U2'].value, 18)
        self.assertEqual(workbook[excel.HISTORY]['AH2'].value, 0)
        self.assertIsNone(workbook[excel.HISTORY]['V2'].value)
        self.assertIsNone(workbook[excel.DETAIL]['K2'].value)
        self.assertIsNone(workbook[excel.DASHBOARD]['B8'].value)
        workbook.close()
        for hole in record['holes'][9:]:
            hole['teeShort'] = False
        record['holes'][17]['teeShort'] = True
        record['updatedAt'] = '2026-11-02T03:00:00.000Z'
        self.write_json([record])
        excel.import_file(self.source, self.output)
        workbook = load_workbook(self.output, data_only=True)
        self.assertEqual(workbook[excel.HISTORY]['AH2'].value, 9)
        self.assertAlmostEqual(workbook[excel.HISTORY]['V2'].value, 1 / 9)
        self.assertAlmostEqual(workbook[excel.DASHBOARD]['B8'].value, 1 / 9)
        workbook.close()

    def test_weighted_rates_and_only_finished_full_rounds_in_dashboard(self):
        one = make_round('one', '2026-11-01')
        for hole in one['holes']:
            hole.update(short=0, direct=True, tee=None)
        one['holes'][0].update(short=1, direct=False, teeShort=True)
        two = make_round('two', '2026-11-02')
        for hole in two['holes'][:9]:
            hole['short'] = 2
        partial = make_round('partial', '2026-11-03')
        partial['holes'][1]['putts'] = None
        unfinished = make_round('unfinished', '2026-11-04')
        unfinished['finished'] = False
        self.write_json([one, two, partial, unfinished])
        excel.import_file(self.source, self.output)
        workbook = load_workbook(self.output, data_only=True)
        dashboard = workbook[excel.DASHBOARD]
        self.assertAlmostEqual(dashboard['B8'].value, 1 / 19)
        self.assertAlmostEqual(dashboard['B9'].value, 9 / 19)
        self.assertEqual(dashboard['A19'].value, '2026-11-01')
        self.assertEqual(dashboard['A20'].value, '2026-11-02')
        self.assertIsNone(dashboard['A21'].value)
        self.assertIsNone(dashboard['C6'].value)
        self.assertIsNone(dashboard['D6'].value)
        self.assertEqual(workbook[excel.HISTORY].max_row, 5)
        workbook.close()

    def test_latest_five_compared_to_previous_five(self):
        records = [make_round(str(index), f'2026-11-{index + 1:02}') for index in range(11)]
        for index, record in enumerate(records):
            record['holes'][0]['long'] = index + 1
        self.write_json(records)
        excel.import_file(self.source, self.output)
        workbook = load_workbook(self.output, data_only=True)
        dashboard = workbook[excel.DASHBOARD]
        self.assertEqual(dashboard['B6'].value, 97)
        self.assertEqual(dashboard['C6'].value, 92)
        self.assertEqual(dashboard['D6'].value, 5)
        workbook.close()

    def test_legacy_id_matches_browser_and_unentered_values_remain_null(self):
        raw = make_round()
        raw['active'] = 2
        raw['view'] = 'result'
        for hole in raw['holes']:
            hole['tee'] = 'good'
        for key in ('id', 'createdAt', 'updatedAt', 'finished'):
            raw.pop(key)
        imported = excel.read_json(raw)[0]
        self.assertEqual(imported['id'], 'legacy-308b0825')
        self.assertEqual(imported['holes'][0]['tee'], 'straight')
        raw['holes'][0]['putts'] = None
        self.assertIsNone(excel.read_json(raw)[0]['holes'][0]['putts'])
        raw['holes'][0]['putts'] = 0
        self.assertEqual(excel.read_json(raw)[0]['holes'][0]['putts'], 0)

    def test_invalid_input_and_unrelated_workbooks_are_not_overwritten(self):
        record = make_round()
        self.write_json([record])
        excel.import_file(self.source, self.output)
        original = self.output.read_bytes()
        for payload in ({'schemaVersion': 99, 'rounds': []}, {'schemaVersion': 2, 'rounds': [record, record]}):
            self.source.write_text(json.dumps(payload), encoding='utf-8')
            with self.assertRaises(ValueError):
                excel.import_file(self.source, self.output)
            self.assertEqual(self.output.read_bytes(), original)
        record['holes'][0]['putts'] = -1
        self.write_json([record])
        with self.assertRaises(ValueError):
            excel.import_file(self.source, self.output)
        other = self.directory / 'unrelated.xlsx'
        wb = Workbook()
        wb.save(other)
        wb.close()
        self.write_json([make_round()])
        with self.assertRaises(ValueError):
            excel.import_file(self.source, other)

    def test_course_names_are_literal_not_excel_formulas(self):
        record = make_round()
        record['place'] = '=HYPERLINK("https://example.com")'
        self.write_json([record])
        excel.import_file(self.source, self.output)
        workbook = load_workbook(self.output, data_only=False)
        self.assertEqual(workbook[excel.HISTORY]['C2'].data_type, 's')
        self.assertEqual(workbook[excel.DETAIL]['C2'].data_type, 's')
        workbook.close()

    def test_equal_timestamp_preserves_existing_round_and_legacy_v1_is_accepted(self):
        record = make_round()
        changed = copy.deepcopy(record)
        changed['holes'][0]['long'] = 5
        self.assertEqual(excel.upsert([record], [changed])[0]['holes'][0]['long'], 2)
        legacy = copy.deepcopy(record)
        legacy['schemaVersion'] = 1
        legacy['holes'][0]['tee'] = 'good'
        imported = excel.read_json(legacy)[0]
        self.assertEqual(imported['id'], record['id'])
        self.assertEqual(imported['holes'][0]['tee'], 'straight')
        first = make_round('first')
        later = make_round('later')
        first['createdAt'] = '2026-11-01T08:00:00+09:00'
        later['createdAt'] = '2026-11-01T00:30:00Z'
        self.assertEqual([r['id'] for r in excel.upsert([], [later, first])], ['first', 'later'])
        changed['updatedAt'] = '2026-11-01T02:00:00.000999Z'
        self.assertEqual(excel.upsert([record], [changed])[0]['holes'][0]['long'], 2)

    @unittest.skipUnless(shutil.which('node'), 'Node.js is needed for cross-language parity')
    def test_statistics_match_browser_core(self):
        record = make_round()
        record['holes'][0].update(short=0, direct=True, tee='left', teeShort=True, obTee=2)
        record['holes'][1].update(short=2, tee=None, teeShort=True, obTee=1, obLater=1)
        record['holes'][2].update(putts=None, tee='right', obTee=1)
        record['holes'][3]['teeShort'] = None
        script = 'const fs=require("fs"),c=require("./roundnote-core.js");process.stdout.write(JSON.stringify(c.analyze(JSON.parse(fs.readFileSync(0,"utf8")))));'
        actual = json.loads(subprocess.check_output(['node', '-e', script], input=json.dumps(record['holes']).encode(), cwd=ROOT))
        expected = excel.analyze(record['holes'])
        for key in ('done', 'score', 'long', 'short', 'putts', 'par', 'teeOB', 'lateOB', 'obTotal', 'bon', 'threePutts', 'approachAttempts', 'approachMisses', 'directCount', 'teeRecorded', 'teeShortRecorded', 'teeShortCount', 'teeDirectionRecorded', 'teeOBDirection'):
            self.assertEqual(expected[key], actual[key], key)
        for key in ('bogeyRate', 'approachRate', 'directRate', 'teeShortRate'):
            self.assertAlmostEqual(expected[key], actual[key] / 100, msg=key)
        legacy = make_round()
        for key in ('id', 'createdAt', 'updatedAt', 'finished'):
            legacy.pop(key)
        for hole in legacy['holes']:
            hole.pop('teeShort')
            hole['tee'] = 'good'
        migration = 'const fs=require("fs"),c=require("./roundnote-core.js");const raw=JSON.parse(fs.readFileSync(0,"utf8"));process.stdout.write(JSON.stringify(c.validateImport(raw).rounds[0]));'
        browser_round = json.loads(subprocess.check_output(['node', '-e', migration], input=json.dumps(legacy).encode(), cwd=ROOT))
        self.assertEqual(excel.read_json(legacy)[0], browser_round)


if __name__ == '__main__':
    unittest.main()
