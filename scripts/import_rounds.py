#!/usr/bin/env python3
"""Import ROUNDNOTE JSON into an updateable, three-sheet Excel workbook.

Usage: python scripts/import_rounds.py roundnote.json --output outputs/roundnote.xlsx
The workbook is rebuilt from canonical round records on each import. Manual workbook
edits are not imported back into ROUNDNOTE; preserve the application's JSON backup.
"""
from __future__ import annotations

import argparse
import json
import math
import os
import re
from pathlib import Path
import tempfile
from datetime import date, datetime, timezone
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.properties import CalcProperties
from openpyxl.utils import get_column_letter

SCHEMA_VERSION = 2
HISTORY = 'ラウンド履歴'
DETAIL = 'ホール別記録'
DASHBOARD = '分析ダッシュボード'
JSON_COLUMN = 52
JSON_HEADER = '_roundnote_json_v2'
TEE = {'left': '左', 'straight': 'まっすぐ', 'right': '右'}
CLUBS = {'driver': 'ドライバー', 'wood_ut': 'ウッド・ユーティリティ', 'iron': 'アイアン', 'wedge': 'ウェッジ'}
APPROACH = {'short': 'ショート', 'over': 'オーバー', 'left': '左', 'right': '右', 'fat': 'ダフリ', 'thin': 'トップ'}
GREEN = '193E2B'
MID_GREEN = '34794E'
MUTED = '65796D'
DEFAULT_GOALS = {'score': 99, 'long': 40, 'ob': 2, 'approachRate': 30, 'putts': 36, 'bogeyOnHoles': 12}


def number(value, maximum=30, nullable=True):
    if value is None and nullable:
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value != int(value) or not 0 <= value <= maximum:
        raise ValueError(f'打数・回数が不正です: {value!r}')
    return int(value)


def iso_date(value):
    if value == '':
        return ''
    if not isinstance(value, str):
        raise ValueError('日付はYYYY-MM-DD形式が必要です')
    try:
        if date.fromisoformat(value).isoformat() != value:
            raise ValueError()
    except ValueError as error:
        raise ValueError(f'日付が不正です: {value!r}') from error
    return value


def timestamp(value):
    if not isinstance(value, str) or re.fullmatch(r'\d{4}-\d{2}-\d{2}T(?:[01]\d|2[0-3]):[0-5]\d:[0-5]\d(?:\.\d{1,9})?(?:Z|[+-](?:[01]\d|2[0-3]):[0-5]\d)', value) is None:
        raise ValueError('作成・更新日時が必要です')
    try:
        parsed = datetime.fromisoformat(value.replace('Z', '+00:00'))
        if parsed.tzinfo is None:
            raise ValueError()
    except ValueError as error:
        raise ValueError(f'更新日時が不正です: {value!r}') from error
    # Match JavaScript Date's millisecond resolution for equal-timestamp upserts.
    return parsed.astimezone(timezone.utc).replace(microsecond=parsed.microsecond // 1000 * 1000)


def normalize_hole(value, legacy=False):
    if not isinstance(value, dict):
        raise ValueError('ホールの記録はオブジェクトが必要です')
    required = ('par', 'long', 'short', 'putts', 'obTee', 'obLater', 'direct')
    if not legacy:
        required += ('tee', 'teeShort', 'approach')
    if any(key not in value for key in required):
        raise ValueError('ホールの記録に必要な項目がありません')
    par = value.get('par', 4)
    if isinstance(par, bool) or par not in (3, 4, 5):
        raise ValueError('PARは3・4・5のみ指定できます')
    tee = value.get('tee')
    if tee == 'good' and legacy:
        tee = 'straight'
    if tee is not None and tee not in TEE:
        raise ValueError(f'ティーショット方向が不正です: {tee!r}')
    club = value.get('teeClub')
    if club is not None and (not isinstance(club, str) or club not in CLUBS):
        raise ValueError('ティーショットのクラブが不正です')
    if 'direct' in value and not isinstance(value['direct'], bool):
        raise ValueError('directは真偽値が必要です')
    if 'teeShort' in value and value['teeShort'] is not None and not isinstance(value['teeShort'], bool):
        raise ValueError('teeShortは真偽値またはnullが必要です')
    approach = value.get('approach')
    if approach is not None:
        if not isinstance(approach, list) or any(not isinstance(k, str) or k not in APPROACH for k in approach):
            raise ValueError('アプローチの記録が不正です')
        if len(set(approach)) != len(approach):
            raise ValueError('アプローチの記録が重複しています')
    short = number(value.get('short'))
    direct = value.get('direct', False)
    if not legacy and direct and short != 0:
        raise ValueError('直接ONと50y以内の打数が一致しません')
    if direct and short is None:
        short = 0
    if short is not None and short > 0:
        direct = False
    return {'par': int(par), 'long': number(value.get('long')), 'short': short,
            'putts': number(value.get('putts')), 'obTee': number(value.get('obTee', 0), 9, False),
            'obLater': number(value.get('obLater', 0), 9, False), 'direct': direct,
            'tee': tee, 'teeShort': value.get('teeShort'), 'approach': approach, 'teeClub': club}


def legacy_id(raw, holes):
    canonical = json.dumps({'date': raw.get('date') or '', 'place': raw.get('place') or '', 'holes': [{k: v for k, v in h.items() if k != 'teeClub'} for h in holes]}, ensure_ascii=False, separators=(',', ':'))
    encoded = canonical.encode('utf-16-le', errors='surrogatepass')
    result = 2166136261
    for pos in range(0, len(encoded), 2):
        result = ((result ^ (encoded[pos] | (encoded[pos + 1] << 8))) * 16777619) & 0xffffffff
    return f'legacy-{result:08x}'


def normalize_round(raw, legacy=False):
    if not isinstance(raw, dict) or not isinstance(raw.get('holes'), list) or len(raw['holes']) != 18:
        raise ValueError('各ラウンドには18ホールの記録が必要です')
    holes = [normalize_hole(h, legacy) for h in raw['holes']]
    day = iso_date(raw.get('date', ''))
    place = raw.get('place', '')
    if not isinstance(place, str) or len(place.encode('utf-16-le', errors='surrogatepass')) // 2 > 100:
        raise ValueError('ゴルフ場名は100文字以内の文字列が必要です')
    active = number(raw.get('active', 0), 17, False)
    identifier = raw.get('id') or (legacy_id(raw, holes) if legacy else None)
    if not isinstance(identifier, str) or re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9._:-]{0,127}', identifier) is None:
        raise ValueError('ラウンドIDが不正です')
    default_time = f'{day or "1970-01-01"}T00:00:00.000Z'
    created = raw.get('createdAt', default_time) if legacy else raw.get('createdAt')
    updated = raw.get('updatedAt', default_time) if legacy else raw.get('updatedAt')
    timestamp(created)
    timestamp(updated)
    if not legacy and any(key not in raw for key in ('active', 'finished', 'createdAt', 'updatedAt')):
        raise ValueError('ラウンドに必要な項目がありません')
    finished = raw.get('finished', False)
    if not isinstance(finished, bool):
        raise ValueError('finishedは真偽値が必要です')
    if legacy:
        finished = finished or all(score(h) is not None for h in holes)
    return {'id': identifier, 'date': day, 'place': place, 'holes': holes, 'active': active,
            'finished': finished, 'createdAt': created, 'updatedAt': updated}


def read_json(data):
    if not isinstance(data, dict):
        raise ValueError('ROUNDNOTEのJSONオブジェクトが必要です')
    version = data.get('schemaVersion')
    if isinstance(version, bool) or version not in (None, 1, 2):
        raise ValueError(f'対応していないスキーマバージョンです: {version!r}')
    if version in (None, 1) and 'holes' in data:
        return [normalize_round(data, legacy=True)]
    raw_rounds = data.get('rounds')
    if not isinstance(raw_rounds, list) or not 1 <= len(raw_rounds) <= 2000:
        raise ValueError('1〜2000件のrounds配列が必要です')
    if version == SCHEMA_VERSION:
        goals = data.get('goals')
        if not isinstance(goals, dict):
            raise ValueError('目標設定がありません')
        for key, lower, upper in [('score', 1, 1620), ('long', 0, 540), ('ob', 0, 162), ('approachRate', 0, 100), ('putts', 0, 540), ('bogeyOnHoles', 0, 18)]:
            n = goals.get(key)
            if isinstance(n, bool) or not isinstance(n, (int, float)) or not math.isfinite(n) or not lower <= n <= upper or (key != 'approachRate' and int(n) != n):
                raise ValueError(f'{key}の目標が不正です')
    rounds = [normalize_round(r, legacy=version != SCHEMA_VERSION) for r in raw_rounds]
    identifiers = [r['id'] for r in rounds]
    if len(set(identifiers)) != len(identifiers):
        raise ValueError('JSON内に重複したラウンドIDがあります')
    if version == SCHEMA_VERSION and data.get('currentRoundId') not in identifiers:
        raise ValueError('現在のラウンドIDが一覧にありません')
    return rounds


def score(hole):
    values = [hole.get(k) for k in ('long', 'short', 'putts')]
    return sum(values) if all(v is not None for v in values) else None


def rate(numerator, denominator):
    return numerator / denominator if denominator else None


def analyze(holes):
    complete = [h for h in holes if score(h) is not None]
    targets = [h for h in complete if not h['direct'] and h['short'] >= 1]
    tee_recorded = [h for h in holes if h['tee'] is not None or h['teeShort'] is True]
    tee_short_recorded = [h for h in tee_recorded if h['teeShort'] is not None]
    tee_directions = [h for h in holes if h['tee'] is not None]
    result = {k: sum(h[k] for h in complete) for k in ('long', 'short', 'putts', 'par')}
    result.update(done=len(complete), score=sum(score(h) for h in complete),
                  teeOB=sum(h['obTee'] for h in holes), lateOB=sum(h['obLater'] for h in holes),
                  bon=sum(h['long'] + h['short'] <= h['par'] - 1 for h in complete),
                  threePutts=sum(h['putts'] >= 3 for h in complete),
                  approachAttempts=len(targets), approachMisses=sum(h['short'] >= 2 for h in targets),
                  directCount=sum(h['direct'] for h in complete), teeRecorded=len(tee_recorded), teeShortRecorded=len(tee_short_recorded),
                  teeDirectionRecorded=len(tee_directions), teeShortCount=sum(h['teeShort'] is True for h in tee_recorded))
    result['obTotal'] = result['teeOB'] + result['lateOB']
    result['teeOBDirection'] = {k: sum(h['obTee'] for h in holes if h['tee'] == k) for k in ('left', 'right')}
    result['teeOBDirection']['other'] = sum(h['obTee'] for h in holes if h['tee'] not in ('left', 'right'))
    for name, num, den in [('bogeyRate', 'bon', 'done'), ('approachRate', 'approachMisses', 'approachAttempts'), ('directRate', 'directCount', 'done'), ('teeShortRate', 'teeShortCount', 'teeShortRecorded')]:
        result[name] = rate(result[num], result[den])
    return result


def previous_rounds(path):
    if not path.exists():
        return []
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if HISTORY not in workbook.sheetnames or workbook[HISTORY].cell(1, JSON_COLUMN).value != JSON_HEADER:
            raise ValueError('出力先はROUNDNOTEで作成したExcelではありません。別の出力名を指定してください')
        records = []
        for row in workbook[HISTORY].iter_rows(min_row=2, min_col=JSON_COLUMN, max_col=JSON_COLUMN):
            if row[0].value:
                records.append(normalize_round(json.loads(row[0].value)))
        if len({r['id'] for r in records}) != len(records):
            raise ValueError('Excel内のラウンドIDが重複しています')
        return records
    finally:
        workbook.close()


def upsert(existing, incoming):
    merged = {r['id']: r for r in existing}
    for record in incoming:
        old = merged.get(record['id'])
        if old is None or timestamp(record['updatedAt']) > timestamp(old['updatedAt']):
            merged[record['id']] = record
    return sorted(merged.values(), key=lambda r: (r['date'], timestamp(r['createdAt']), r['id']))


class Writer:
    def __init__(self):
        self.workbook = Workbook()
        self.workbook.remove(self.workbook.active)
        self.cached = {}
        for name in (DASHBOARD, HISTORY, DETAIL):
            sheet = self.workbook.create_sheet(name)
            sheet.sheet_view.showGridLines = False
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.sheet_properties.tabColor = GREEN if name == DASHBOARD else None
            self.cached[name] = {}
        self.workbook.calculation = CalcProperties(fullCalcOnLoad=True, forceFullCalc=True)

    def value(self, sheet, row, column, value):
        cell = sheet.cell(row, column, value)
        if isinstance(value, str) and value.startswith('='):
            cell.data_type = 's'
        return cell

    def formula(self, sheet, row, column, expression, cached):
        sheet.cell(row, column, expression)
        self.cached[sheet.title][f'{get_column_letter(column)}{row}'] = cached

    def table(self, sheet, headers, name, header_row=1):
        for col, header in enumerate(headers, 1):
            cell = self.value(sheet, header_row, col, header)
            cell.fill = PatternFill('solid', fgColor=GREEN)
            cell.font = Font(name='Arial', size=10, color='FFFFFF', bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        sheet.row_dimensions[header_row].height = 32
        last = sheet.max_row
        if last > header_row:
            table = Table(displayName=name, ref=f'A{header_row}:{get_column_letter(len(headers))}{last}')
            table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium4', showRowStripes=True, showColumnStripes=False)
            sheet.add_table(table)
        sheet.auto_filter.ref = f'A{header_row}:{get_column_letter(len(headers))}{max(header_row, last)}'
        sheet.freeze_panes = 'D2'
        for row in sheet.iter_rows(min_row=header_row + 1, max_row=last, max_col=len(headers)):
            for cell in row:
                cell.font = Font(name='Arial', size=10, color=GREEN)
                cell.alignment = Alignment(vertical='center')
        for col in range(1, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 16
        sheet.column_dimensions['A'].width = 28
        sheet.column_dimensions['B'].width = 13
        sheet.column_dimensions['C'].width = 28
        sheet.print_options.horizontalCentered = True


DETAIL_HEADERS = ['ラウンドID', '日付', 'ゴルフ場', 'ホール', '前後半', 'PAR', '50yまで', '50yからON', 'パット', 'ティー方向', 'ティーショート', '直接ON', 'ティーOB', '2打目以降OB', '50yミス', '入力完了', 'スコア', 'ONまで', 'ボギーオン以内', '50y対象', '50y2打以上', '直接ON対象', 'ティー結果記録', 'ティー方向記録', 'ティーOB左', 'ティーOB右', 'ティーOBその他', '3パット以上', 'ティーショート対象', 'ティー使用クラブ']
HISTORY_HEADERS = ['ラウンドID', '日付', 'ゴルフ場', '終了', '更新日時', '入力済みH', '終了済み18H', 'スコア', '対PAR', '50yまで', '50yからON', 'パット', 'OB合計', 'ティーOB', '2打目以降OB', 'ティーOB左', 'ティーOB右', 'ティーOBその他', 'ティー方向記録H', 'ティーショートH', 'ティー結果記録H', 'ティーショート率', '50y対象H', '50y2打以上H', '50y2打以上率', 'ボギーオン以内H', 'ボギーオン率', '直接ON回数', '直接ON率', '3パット以上H', 'OUTスコア', 'INスコア', '作成日時', 'ティーショート対象H']


def add_detail(writer, record):
    sheet = writer.workbook[DETAIL]
    first = sheet.max_row + 1 if sheet.cell(1, 1).value is not None else 2
    for index, h in enumerate(record['holes']):
        row = first + index
        raw = [record['id'], date.fromisoformat(record['date']) if record['date'] else None, record['place'], index + 1, 'OUT' if index < 9 else 'IN', h['par'], h['long'], h['short'], h['putts'], TEE.get(h['tee']), None if h['teeShort'] is None else int(h['teeShort']), int(h['direct']), h['obTee'], h['obLater'], None if h['approach'] is None else '・'.join(APPROACH[k] for k in h['approach']) if h['approach'] else 'ミスなし']
        for col, value in enumerate(raw, 1):
            writer.value(sheet, row, col, value)
        writer.value(sheet, row, 30, CLUBS.get(h.get('teeClub')))
        complete = score(h) is not None
        on = h['long'] + h['short'] if h['long'] is not None and h['short'] is not None else None
        target = complete and not h['direct'] and h['short'] >= 1
        calculations = [
            (16, f'=--AND(ISNUMBER(G{row}),ISNUMBER(H{row}),ISNUMBER(I{row}))', int(complete)),
            (17, f'=IF(P{row},SUM(G{row}:I{row}),"")', score(h)),
            (18, f'=IF(AND(ISNUMBER(G{row}),ISNUMBER(H{row})),G{row}+H{row},"")', on),
            (19, f'=IF(P{row},--(R{row}<=F{row}-1),"")', int(on <= h['par'] - 1) if complete else None),
            (20, f'=IF(P{row},--AND(NOT(L{row}),H{row}>=1),"")', int(target) if complete else None),
            (21, f'=IF(P{row},--AND(T{row}=1,H{row}>=2),"")', int(target and h['short'] >= 2) if complete else None),
            (22, f'=IF(P{row},--L{row},"")', int(h['direct']) if complete else None),
            (23, f'=--OR(J{row}<>"",K{row}=1)', int(h['tee'] is not None or h['teeShort'] is True)),
            (24, f'=--(J{row}<>"")', int(h['tee'] is not None)),
            (25, f'=IF(J{row}="左",M{row},0)', h['obTee'] if h['tee'] == 'left' else 0),
            (26, f'=IF(J{row}="右",M{row},0)', h['obTee'] if h['tee'] == 'right' else 0),
            (27, f'=IF(AND(J{row}<>"左",J{row}<>"右"),M{row},0)', h['obTee'] if h['tee'] not in ('left', 'right') else 0),
            (28, f'=IF(P{row},--(I{row}>=3),"")', int(h['putts'] >= 3) if complete else None),
            (29, f'=--AND(ISNUMBER(K{row}),OR(J{row}<>"",K{row}=1))', int(h['teeShort'] is not None and (h['tee'] is not None or h['teeShort'] is True))),
        ]
        for col, expression, cached in calculations:
            writer.formula(sheet, row, col, expression, cached)
        sheet.cell(row, 2).number_format = 'yyyy-mm-dd'
    return first, first + 17


def add_history(writer, record, first, last):
    sheet = writer.workbook[HISTORY]
    row = sheet.max_row + 1 if sheet.cell(1, 1).value is not None else 2
    t = analyze(record['holes'])
    r = lambda col: f"'{DETAIL}'!${col}${first}:${col}${last}"
    raw = [record['id'], date.fromisoformat(record['date']) if record['date'] else None, record['place'], record['finished'], record['updatedAt']]
    for col, value in enumerate(raw, 1):
        writer.value(sheet, row, col, value)
    calculations = [
        (6, f'=COUNTIFS({r("P")},1)', t['done']),
        (7, f'=AND(D{row},F{row}=18)', record['finished'] and t['done'] == 18),
        (8, f'=IF(F{row}=0,"",SUM({r("Q")}))', t['score'] if t['done'] else None),
        (9, f'=IF(F{row}=0,"",H{row}-SUMIFS({r("F")},{r("P")},1))', t['score'] - t['par'] if t['done'] else None),
        (10, f'=IF(F{row}=0,"",SUMIFS({r("G")},{r("P")},1))', t['long'] if t['done'] else None),
        (11, f'=IF(F{row}=0,"",SUMIFS({r("H")},{r("P")},1))', t['short'] if t['done'] else None),
        (12, f'=IF(F{row}=0,"",SUMIFS({r("I")},{r("P")},1))', t['putts'] if t['done'] else None),
        (13, f'=SUM(N{row}:O{row})', t['obTotal']),
        (14, f'=SUM({r("M")})', t['teeOB']), (15, f'=SUM({r("N")})', t['lateOB']),
        (16, f'=SUM({r("Y")})', t['teeOBDirection']['left']), (17, f'=SUM({r("Z")})', t['teeOBDirection']['right']), (18, f'=SUM({r("AA")})', t['teeOBDirection']['other']),
        (19, f'=SUM({r("X")})', t['teeDirectionRecorded']),
        (20, f'=COUNTIFS({r("W")},1,{r("K")},1)', t['teeShortCount']),
        (21, f'=SUM({r("W")})', t['teeRecorded']),
        (22, f'=IF(AH{row}=0,"",T{row}/AH{row})', t['teeShortRate']),
        (23, f'=SUM({r("T")})', t['approachAttempts']), (24, f'=SUM({r("U")})', t['approachMisses']),
        (25, f'=IF(W{row}=0,"",X{row}/W{row})', t['approachRate']),
        (26, f'=SUM({r("S")})', t['bon']), (27, f'=IF(F{row}=0,"",Z{row}/F{row})', t['bogeyRate']),
        (28, f'=SUM({r("V")})', t['directCount']), (29, f'=IF(F{row}=0,"",AB{row}/F{row})', t['directRate']),
        (30, f'=SUM({r("AB")})', t['threePutts']),
        (34, f'=SUM({r("AC")})', t['teeShortRecorded']),
    ]
    for col, half in ((31, 'OUT'), (32, 'IN')):
        partial = analyze(record['holes'][:9] if half == 'OUT' else record['holes'][9:])
        calculations.append((col, f'=IF(COUNTIFS({r("E")},"{half}",{r("P")},1)=0,"",SUMIFS({r("Q")},{r("E")},"{half}"))', partial['score'] if partial['done'] else None))
    for col, expression, cached in calculations:
        writer.formula(sheet, row, col, expression, cached)
    writer.value(sheet, row, 33, record['createdAt'])
    writer.value(sheet, row, JSON_COLUMN, json.dumps(record, ensure_ascii=False, separators=(',', ':')))
    sheet.cell(row, 2).number_format = 'yyyy-mm-dd'
    for col in (22, 25, 27, 29):
        sheet.cell(row, col).number_format = '0.0%'
    return row, t


def period_values(items):
    if not items:
        return [None] * 6
    metrics = [t for _, t, _ in items]
    return [sum(m['score'] for m in metrics) / len(metrics), sum(m['obTotal'] for m in metrics) / len(metrics),
            rate(sum(m['teeShortCount'] for m in metrics), sum(m['teeShortRecorded'] for m in metrics)),
            rate(sum(m['approachMisses'] for m in metrics), sum(m['approachAttempts'] for m in metrics)),
            sum(m['putts'] for m in metrics) / len(metrics), rate(sum(m['bon'] for m in metrics), sum(m['done'] for m in metrics))]


def add_dashboard(writer, records):
    sheet = writer.workbook[DASHBOARD]
    sheet.column_dimensions['A'].width = 24
    for col in range(2, 13):
        sheet.column_dimensions[get_column_letter(col)].width = 15
    writer.value(sheet, 2, 1, 'ROUNDNOTE 分析')
    sheet['A2'].font = Font(name='Arial', size=16, color=GREEN, bold=True)
    for col, label in enumerate(['指標', '直近5R', '前5R', '変化'], 1):
        writer.value(sheet, 5, col, label)
    eligible = [(r, t, record) for r, t, record in records if record['finished'] and t['done'] == 18]
    recent, before = eligible[-5:], eligible[-10:-5]
    newest, older = period_values(recent), period_values(before)
    labels = ['総スコア（打）', 'OB（回）', 'ティーショート率', '50yから2打以上率', 'パット（打）', 'ボギーオン以内率']
    count_columns = ['H', 'M', None, None, 'L', None]
    ratio_columns = {2: ('T', 'AH'), 3: ('X', 'W'), 5: ('Z', 'F')}
    for index, label in enumerate(labels):
        row = 6 + index
        writer.value(sheet, row, 1, label)
        for col, period, cached_values in ((2, recent, newest), (3, before, older)):
            if not period:
                continue
            refs = [f"'{HISTORY}'!{count_columns[index]}{r}" for r, _, _ in period] if count_columns[index] else None
            if refs:
                expression = f'=AVERAGE({",".join(refs)})'
            else:
                numerator, denominator = ratio_columns[index]
                num = ','.join(f"'{HISTORY}'!{numerator}{r}" for r, _, _ in period)
                den = ','.join(f"'{HISTORY}'!{denominator}{r}" for r, _, _ in period)
                expression = f'=IF(SUM({den})=0,"",SUM({num})/SUM({den}))'
            writer.formula(sheet, row, col, expression, cached_values[index])
        cached_delta = newest[index] - older[index] if newest[index] is not None and older[index] is not None else None
        if cached_delta is not None and index in ratio_columns:
            cached_delta *= 100
        multiplier = '*100' if index in ratio_columns else ''
        writer.formula(sheet, row, 4, f'=IF(AND(ISNUMBER(B{row}),ISNUMBER(C{row})),(B{row}-C{row}){multiplier},"")', cached_delta)
        for col in (2, 3):
            sheet.cell(row, col).number_format = '0.0%' if index in ratio_columns else '0.0'
        sheet.cell(row, 4).number_format = '+0.0"pt";-0.0"pt";0.0"pt"' if index in ratio_columns else '+0.0;-0.0;0.0'
    writer.value(sheet, 13, 1, f'終了済み18Hのみ。直近 {len(recent)}R、前 {len(before)}R。率は対象数で加重集計。変化は直近−前。')
    sheet['A13'].font = Font(name='Arial', size=10, color=MUTED)
    writer.value(sheet, 14, 1, '対象がない率は空欄。JSON再取り込みで更新。Excelでの手修正は次回更新で置き換わります。')
    sheet['A14'].font = Font(name='Arial', size=10, color=MUTED)
    headers = ['日付', 'ゴルフ場', 'スコア', 'OB合計', 'ティーOB左', 'ティーOB右', 'ティーOBその他', '2打目以降OB', 'ティーショート率', '50y2打以上率', 'パット', 'ボギーオン率']
    for col, header in enumerate(headers, 1):
        writer.value(sheet, 18, col, header)
    mappings = [('H', 'score'), ('M', 'obTotal'), ('P', None), ('Q', None), ('R', None), ('O', 'lateOB'), ('V', 'teeShortRate'), ('Y', 'approachRate'), ('L', 'putts'), ('AA', 'bogeyRate')]
    for row, (history_row, t, record) in enumerate(eligible, 19):
        writer.value(sheet, row, 1, record['date'])
        writer.value(sheet, row, 2, record['place'])
        for col, (source_col, metric) in enumerate(mappings, 3):
            value = t[metric] if metric else t['teeOBDirection'][{5: 'left', 6: 'right', 7: 'other'}[col]]
            writer.formula(sheet, row, col, f'=\'{HISTORY}\'!{source_col}{history_row}', value)
            if col in (9, 10, 12):
                sheet.cell(row, col).number_format = '0.0%'
    if not eligible:
        writer.value(sheet, 20, 1, '終了済み18Hの記録がありません')
    chart_start = max(23, 21 + len(eligible))
    specs = [('スコア推移（打）', [3], 'A', 0, False), ('OB内訳（回）', [5, 6, 7, 8], 'G', 0, False), ('ティーショート率', [9], 'A', 16, True), ('50y以内から2打以上率', [10], 'G', 16, True), ('パット数（打）', [11], 'A', 32, False), ('ボギーオン以内率', [12], 'G', 32, True)]
    if eligible:
        for title, columns, anchor, offset, percent in specs:
            chart = LineChart()
            chart.title = title
            chart.style = 13
            chart.width = 16.2
            chart.height = 8
            chart.display_blanks = 'gap'
            chart.legend = None if len(columns) == 1 else chart.legend
            if chart.legend is not None:
                chart.legend.position = 'b'
            chart.y_axis.numFmt = '0%' if percent else '0'
            chart.y_axis.scaling.min = 0 if percent or len(columns) > 1 else None
            if len(columns) > 1:
                chart.y_axis.majorUnit = 1
            if percent:
                chart.y_axis.scaling.max = 1
            chart.x_axis.tickLblPos = 'low'
            for col in columns:
                chart.add_data(Reference(sheet, min_col=col, min_row=18, max_row=18 + len(eligible)), titles_from_data=True)
            chart.set_categories(Reference(sheet, min_col=1, min_row=19, max_row=18 + len(eligible)))
            for series, col, color in zip(chart.series, columns, ['34794E', '9B6743', '839583', '416884']):
                series.title = SeriesLabel(v=sheet.cell(18, col).value)
                series.graphicalProperties.line.solidFill = color
                series.graphicalProperties.line.width = 22000
                series.smooth = False
            sheet.add_chart(chart, f'{anchor}{chart_start + offset}')
    for row in sheet.iter_rows(min_row=5, max_row=max(20, 18 + len(eligible)), max_col=12):
        for cell in row:
            if cell.row not in (13, 14):
                cell.font = Font(name='Arial', size=10, color=GREEN)
            cell.alignment = Alignment(vertical='center')
    for row_number in (5, 18):
        for col in range(1, 5 if row_number == 5 else 13):
            cell = sheet.cell(row_number, col)
            cell.fill = PatternFill('solid', fgColor=GREEN)
            cell.font = Font(name='Arial', size=10, color='FFFFFF', bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        sheet.row_dimensions[row_number].height = 32
    for row in range(6, 12):
        sheet.row_dimensions[row].height = 24
        good = 'greaterThan' if row == 11 else 'lessThan'
        bad = 'lessThan' if row == 11 else 'greaterThan'
        sheet.conditional_formatting.add(f'D{row}', CellIsRule(operator=good, formula=['0'], font=Font(color=MID_GREEN)))
        sheet.conditional_formatting.add(f'D{row}', CellIsRule(operator=bad, formula=['0'], font=Font(color='9B6743')))
    sheet.print_area = f'A1:L{chart_start + 47 if eligible else 22}'
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0


def write_cached_values(path, cached):
    """Keep initial Excel/chart values available before Excel's first recalculation."""
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    ET.register_namespace('', ns['s'])
    with zipfile.ZipFile(path) as archive:
        files = {item.filename: (item, archive.read(item.filename)) for item in archive.infolist()}
    for index, (_, cells) in enumerate(cached.items(), 1):
        filename = f'xl/worksheets/sheet{index}.xml'
        item, body = files[filename]
        root = ET.fromstring(body)
        for cell in root.findall('.//s:c', ns):
            address = cell.attrib.get('r')
            if address not in cells:
                continue
            value = cells[address]
            element = cell.find('s:v', ns)
            if element is None:
                element = ET.SubElement(cell, f'{{{ns["s"]}}}v')
            if value is None:
                cell.set('t', 'str')
                element.text = ''
            elif isinstance(value, bool):
                cell.set('t', 'b')
                element.text = '1' if value else '0'
            else:
                cell.attrib.pop('t', None)
                element.text = str(value)
        files[filename] = (item, ET.tostring(root, encoding='utf-8', xml_declaration=True))
    with zipfile.ZipFile(path, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        for item, body in files.values():
            archive.writestr(item, body)


def write_workbook(rounds, output):
    writer = Writer()
    records = []
    # Reserve the headers before rows so all 18-hole blocks have stable ranges.
    for name, headers in ((HISTORY, HISTORY_HEADERS), (DETAIL, DETAIL_HEADERS)):
        for col, header in enumerate(headers, 1):
            writer.value(writer.workbook[name], 1, col, header)
    writer.value(writer.workbook[HISTORY], 1, JSON_COLUMN, JSON_HEADER)
    for record in rounds:
        first, last = add_detail(writer, record)
        row, metrics = add_history(writer, record, first, last)
        records.append((row, metrics, record))
    writer.table(writer.workbook[HISTORY], HISTORY_HEADERS, 'RoundHistory')
    writer.table(writer.workbook[DETAIL], DETAIL_HEADERS, 'HoleRecords')
    writer.workbook[HISTORY].column_dimensions['AZ'].hidden = True
    writer.workbook[HISTORY].column_dimensions['E'].width = 26
    writer.workbook[HISTORY].column_dimensions['AG'].width = 26
    writer.workbook[DETAIL].column_dimensions['O'].width = 28
    writer.workbook[DETAIL].column_dimensions['AD'].width = 28
    add_dashboard(writer, records)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(prefix='roundnote-', suffix='.xlsx', dir=output.parent)
    os.close(descriptor)
    try:
        writer.workbook.save(temporary)
        write_cached_values(temporary, writer.cached)
        os.replace(temporary, output)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def import_file(source, output):
    source, output = Path(source), Path(output)
    if output.suffix.lower() != '.xlsx':
        raise ValueError('出力先には.xlsxファイルを指定してください')
    with source.open(encoding='utf-8-sig') as handle:
        incoming = read_json(json.load(handle))
    merged = upsert(previous_rounds(output), incoming)
    write_workbook(merged, output)
    return len(merged)


def main(argv=None):
    parser = argparse.ArgumentParser(description='ROUNDNOTEのJSONをExcelに取り込みます。同じIDは更新し、古い更新日時の記録は上書きしません。')
    parser.add_argument('json_file', type=Path)
    parser.add_argument('--output', type=Path, default=Path('outputs/roundnote.xlsx'))
    args = parser.parse_args(argv)
    try:
        total = import_file(args.json_file, args.output)
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as error:
        parser.exit(1, f'取り込みできません: {error}\n')
    print(f'{total}ラウンドを保存しました: {args.output}')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
